# -*- coding: utf-8 -*-
"""美的/容声报价 + 表分发 + 合卡（2026-10-11 拆分一期二级细拆：从 main.py 原样搬出，纯搬运零行为变化）。
load_records_from_path（认表+分发）/ merge_records（四批合并成卡）/ 补贴表解析。
依赖方向：normalize → parsing（本文件是 parsing 域的收口，被 payload 调用）。"""
import datetime as _dt
import io
import json
import os
import re

import openpyxl  # ⚠ 顶层 import（PyInstaller 规矩：不能放函数体内）

from normalize import brand_norm, core_of, norm_model, num, sx
from paths import rp
from state import _tick
from xlsx_core import col_of, header_map, nc, rows_of, sc
from parsers_haier import parse_clean, parse_price

# 「美的」sheet 是自己维护的固定格式：r1 表头、数据 r2+。列序写死并做校验 ——
# 不能走 header_map，因为表头里「0.016」会被 _is_placeholder_label 当成日期序列号
# 占位符，然后被下一行的 69.5999 顶掉。
MD_LAYOUT = {"cat": 1, "brand": 2, "code": 3, "name": 4, "mo": 5,
             "mh": 6, "m2": 7, "mr": 8, "lp": 9, "jd": 11,
             "note": 13, "xy": 17, "profit": 18}


def parse_meidi(ws):
    """「美的」sheet → 记录列表（kind='m'）。另一套报价体系：
    拿货价（慧买卖）/ 价格（给二批）/ 价格取整 / 零售指导价 / 京东价 / 参考闲鱼价 / 利润。"""
    rows = rows_of(ws)
    if not rows:
        return []
    hdr = [sx(x) for x in (rows[0] or [])]
    i_mo = MD_LAYOUT["mo"]
    if len(hdr) <= max(MD_LAYOUT.values()) or hdr[i_mo] != "型号":
        raise RuntimeError("「美的」sheet 列序变了，请核对表头：%s" % (hdr[:22],))

    def cell(r, k):
        i = MD_LAYOUT[k]
        return r[i] if i < len(r) else None

    out = []
    for r in rows[1:]:
        if not r:
            continue
        mo = sx(cell(r, "mo"))
        if not mo or norm_model(mo) == "型号":
            continue
        out.append({
            "kind": "m", "mo": mo,
            "b": brand_norm(sx(cell(r, "brand")), sx(cell(r, "name"))),
            "cat": sx(cell(r, "cat")),
            "code": sx(cell(r, "code")),
            "n": sx(cell(r, "name")),
            "lp": num(cell(r, "lp")),
            "mh": num(cell(r, "mh")),          # 拿货价（慧买卖）
            "m2": num(cell(r, "m2")),          # 价格（给二批）
            "mr": num(cell(r, "mr")),          # 价格取整
            "jd": num(cell(r, "jd")),          # 京东价或国补后价
            "xyref": num(cell(r, "xy")),       # 参考闲鱼价
            "profit": num(cell(r, "profit")),  # 按参考闲鱼价算的利润
            "note": sx(cell(r, "note")),
        })
    return out


def find_sheet(names, pred):
    for n in names:
        if pred(n):
            return n
    return None


def pick_sheet(names, *keys):
    for k in keys:
        for n in names:
            if k in n:
                return n
    return None


def load_records_from_path(path):
    """读一个 xlsx → (记录列表, 月份提示)。自动识别是哪个表。"""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    names = wb.sheetnames
    base = os.path.basename(path)
    month = ""

    if "滞销" in base or find_sheet(names, lambda n: "滞销清理明细" in n):
        nl = pick_sheet(names, "滞销清理明细", "滞销")
        ns = pick_sheet(names, "样机清理明细", "样机")
        recs = []
        if nl:
            recs += parse_clean(wb[nl], "l")
            m = re.search(r"(\d{1,2})\s*月", nl)
            if m:
                month = "%d-%02d" % (_dt.datetime.now().year, int(m.group(1)))
        if ns:
            recs += parse_clean(wb[ns], "s")
            if not month:
                m = re.search(r"(\d{1,2})\s*月", ns)
                if m:
                    month = "%d-%02d" % (_dt.datetime.now().year, int(m.group(1)))
        return recs, month

    np_ = pick_sheet(names, "型号价格表", "价格表", "只看这个")
    if np_:
        recs = parse_price(wb[np_])
        # 「美的」页：另一套报价体系（美的大家电 / 美的小家电 / 容声）
        if "美的" in names:
            mrecs = parse_meidi(wb["美的"])
            print("  美的 %d 条" % len(mrecs))
            recs += mrecs
        # 「现存量10.6」不再解析（2026-10 改版：库存/仓库全部走用友实时接口）
        return recs, ""

    raise RuntimeError("认不出这个文件是哪个表：" + base)


def merge_records(allrecs):
    """公布版（滞销/样机）+ 价格表 → 按 core 合并成一张卡。"""
    rows = {}

    # ---- 第一批：公布版清理明细
    for rec in [x for x in allrecs if x.get("kind") in ("l", "s")]:
        ck = core_of(rec["mo"])
        r = rows.setdefault(ck, {"c": ck, "mo": rec["mo"], "b": rec["b"], "cat": rec["cat"],
                                 "code": rec["code"], "n": rec["n"], "lp": rec["lp"],
                                 "nat": rec["nat"], "natk": rec.get("natk"), "k": "",
                                 "src_kind": "c"})
        for k in ("b", "cat", "code", "n", "lp"):
            if r.get(k) in (None, "") and rec.get(k) not in (None, ""):
                r[k] = rec[k]
        if rec["kind"] == "l":
            if "l" not in r["k"]:
                r["k"] += "l"
            r.update(lq=rec["q"], lpz=rec["p"], lb=rec["rb"], lav=rec["av"])
        else:
            if "s" not in r["k"]:
                r["k"] += "s"
            r.update(sq=rec["q"], spz=rec["p"], sb=rec["rb"], sav=rec["av"], st=rec["st"])
        if rec.get("mo") and len(str(rec["mo"])) > len(str(r.get("mo") or "")):
            r["mo"] = rec["mo"]
        if rec.get("nat") is not None:
            r["nat"], r["natk"] = rec["nat"], rec.get("natk")

    # ---- 第二批：价格表
    #   精准归卡（2026-10-10 两轮定稿）：
    #   ① 先按「型号原文」分组 —— ES60H-D2S 和 ES60H-D2S新 编码/价格都不同，是两个 SKU，
    #      不能因 core_of 截中文并成一张卡互覆价格（D2S新 的行把 D2S 的隆阳价盖掉过）；
    #   ② 同一型号原文有多个货源 → 合并成一张卡（2026-10-10 晚用户拍板：不拆卡），
    #      卡上挂 srcs 列表，每个货源各标各的价，前端价格格/发货行注明货源；
    #      卡主字段取「有物料编码」的货源行（编码是库存/图/官网的主键，隆阳这类正常供货商都有；
    #      HM夏那种裸型号行没编码，只当附加价）——库存口径跟着主货源走；
    #   ③ 同货源多行 → 沿用老口径取最后一行；
    #   ④ 公布版先建的卡只跟「型号原文一致」的价格表行合并，原文不同的自己开卡。
    #   发布粒度（用户拍板）：一型一条发布，货源只是成本口径标注，发布键不带货源。
    _px = {}
    for rec in [x for x in allrecs if x.get("kind") is None]:
        _px.setdefault(core_of(rec["mo"]), {}).setdefault(norm_model(rec["mo"]), []).append(rec)
    for ck_core, groups in _px.items():
        multi_model = len(groups) > 1
        base0 = rows.get(ck_core)                                  # 公布版先建的卡（可能没有）
        base_ex = norm_model(base0["mo"]) if base0 else None
        for ex, recs in groups.items():
            by_src = {}
            for rec in recs:                                       # dict 保序：货源按首次出现排
                by_src.setdefault(rec.get("from_src") or "", []).append(rec)
            src_rows = []                                          # [(货源名, 该货源代表行)]
            for sname, srecs in by_src.items():
                src_rows.append((sname, srecs[-1]))                # 同货源多行取最后一行（老口径）
            pri = next((r for _s, r in src_rows if r.get("code")), src_rows[0][1])
            rec = pri
            ck = ck_core
            if multi_model and base_ex != ex:
                ck = "%s@%s" % (ck_core, str(recs[0].get("mo") or ex))
            r = rows.get(ck)
            if r is None:
                r = rows[ck] = {"c": ck, "mo": rec["mo"], "b": rec["b"], "cat": rec["cat"],
                                "code": rec["code"], "n": rec["n"], "lp": rec["lp"],
                                "nat": None, "natk": None, "k": "", "src_kind": "p"}
            else:
                r["src_kind"] = "cp"
            for k in ("b", "cat", "code", "n", "lp"):
                if r.get(k) in (None, "") and rec.get(k) not in (None, ""):
                    r[k] = rec[k]
            r.update(pt=rec["pt"],
                     ship=rec["ship"], sfrom=rec["from_src"] or r.get("sfrom"),
                     xyid=rec["xyid"], xyp=rec["xyp"], profit=rec["profit"],
                     s69=rec["s69"], smax=rec["smax"], searn=rec["searn"], suser=rec["suser"],
                     base=rec["base"], t2=rec["t2"])
            if rec.get("mo") and len(str(rec["mo"])) > len(str(r.get("mo") or "")):
                r["mo"] = rec["mo"]
            # 多货源 → 挂 srcs（含主货源在内的每个货源一份价），前端据此注明「这个价是哪个货源的」
            if len(src_rows) > 1:
                r["srcs"] = [{"sfrom": s, "pt": sr.get("pt"), "lpz": sr.get("lpz"),
                              "spz": sr.get("spz"), "ship": sr.get("ship"),
                              "s69": sr.get("s69"), "xyp": sr.get("xyp")}
                             for s, sr in src_rows]

    # ---- 第三批：美的 / 容声（自己那一套报价）
    #   同一型号在这张表里可能出现两次（不同批次报价），取「拿货价·慧买卖」最低的那行；
    #   整行一起换，免得出现「拿货价取的是 A 行、参考闲鱼价取的是 B 行」的错配。
    def _mkey(x):
        return x["mh"] if isinstance(x.get("mh"), (int, float)) else float("inf")
    mdedup = {}
    for rec in [x for x in allrecs if x.get("kind") == "m"]:
        ck = core_of(rec["mo"])
        cur = mdedup.get(ck)
        if cur is None or _mkey(rec) < _mkey(cur):
            mdedup[ck] = rec
    for ck, rec in mdedup.items():
        r = rows.get(ck)
        if r is None:
            r = rows[ck] = {"c": ck, "mo": rec["mo"], "b": rec["b"], "cat": rec["cat"],
                            "code": rec["code"], "n": rec["n"], "lp": rec["lp"],
                            "nat": None, "natk": None, "k": "", "src_kind": "m"}
        else:
            # 主表里也有这个型号 → 只补它没有的字段，别覆盖
            if "m" not in (r.get("src_kind") or ""):
                r["src_kind"] = (r.get("src_kind") or "") + "m"
        for k in ("b", "cat", "code", "n", "lp"):
            if r.get(k) in (None, "") and rec.get(k) not in (None, ""):
                r[k] = rec[k]
        for k in ("mh", "m2", "mr", "jd", "xyref", "profit", "note"):
            if r.get(k) in (None, "") and rec.get(k) not in (None, ""):
                r[k] = rec[k]
        if rec.get("mo") and len(str(rec["mo"])) > len(str(r.get("mo") or "")):
            r["mo"] = rec["mo"]

    # ---- 第四批：官网价快照（打包进 assets 的 midea_prices.json / ronshen_prices.json，
    #   分别抓自美的商城、海信商城，构建期随包更新）。缓存键 = 合并后 rows 的 mo。
    mcache = {}
    for fname in ("midea_prices.json", "ronshen_prices.json"):
        try:
            fp = rp(fname)
            if os.path.isfile(fp):
                with io.open(fp, encoding="utf-8") as f:
                    mcache.update(json.load(f))
        except Exception:
            pass
    nmprice = 0
    for ck, r in rows.items():
        p = mcache.get(r.get("mo"))
        if p and p.get("price") is not None:
            r["mp"] = p["price"]
            if p.get("oprice") is not None:
                r["mop"] = p["oprice"]
            nmprice += 1
    _tick("官网价合并 %d/%d" % (nmprice, len(mcache)))

    # ---- 品牌归一（修「品牌列填成品类」的笔误 + 美的/容声合并）
    for r in rows.values():
        r["b"] = brand_norm(r.get("b"), r.get("n"))
    return rows


# ------------------------------------------------------------------ 补贴表解析（深圳以旧换新集团商品中心）


def _norm_txt(v):
    return str(v or "").strip()


def parse_subsidy_rows(xlsx_path):
    """以旧换新表 → [{mo(归一), c(69码), p(零售价)}]。"""
    wb = openpyxl.load_workbook(xlsx_path, read_only=True, data_only=True)
    out = []
    try:
        for ws in wb.worksheets:
            rows = []
            for row in ws.iter_rows(values_only=True):
                if any(_norm_txt(c) for c in (row or [])):
                    rows.append(list(row))
                # 83303 行的大表，全量读入（read_only 模式内存可控）
            if not rows:
                continue
            # 找表头行：前 15 行里同时含「型号」「条码/编码」的那行
            # 实表结构：1=商品条码(69码) 2=商品编码(助记码) 10=型号 20=售价
            hi, i_mo, i_code, i_price = None, None, None, None
            for ri, row in enumerate(rows[:15]):
                cells = [_norm_txt(c) for c in row]
                if "型号" in cells and ("商品条码" in cells or "商品编码" in cells):
                    hi = ri
                    for ci, h in enumerate(cells):
                        if h == "型号" and i_mo is None:
                            i_mo = ci
                        elif h == "商品条码" and i_code is None:
                            i_code = ci                     # 69码优先取条码列
                        elif h == "商品编码" and i_code is None:
                            i_code = ci
                        elif (h == "售价" or h == "零售价") and i_price is None:
                            i_price = ci
                    break
            if hi is None or i_mo is None or i_code is None:
                continue
            if i_price is None:
                i_price = i_code + 1 if i_code + 1 < len(rows[hi]) else None
            n = 0
            for row in rows[hi + 1:]:
                mo = norm_model(row[i_mo] if i_mo < len(row) else None)
                code = _norm_txt(row[i_code] if i_code is not None and i_code < len(row) else None)
                if not mo or not code:
                    continue
                pr = None
                if i_price is not None and i_price < len(row):
                    v = row[i_price]
                    if isinstance(v, (int, float)):
                        pr = float(v)
                    else:
                        v2 = re.sub(r"[^\d.]", "", str(v or ""))
                        if v2:
                            try:
                                pr = float(v2)
                            except ValueError:
                                pr = None
                out.append({"mo": mo, "c": code, "p": pr})
                n += 1
            if n:
                break        # 第一个有数据的 sheet 就够
    finally:
        wb.close()
    return out
